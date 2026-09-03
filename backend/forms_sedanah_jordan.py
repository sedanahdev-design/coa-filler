"""Fill logic for the 'Sedanah Jordan' form (invoice, packing list).

Both documents share the same header table: 'From:' is static (Sedanah's own
info); Port of Loading/Discharge/Final Destination, Country of Final
Destination and Carriage are taken from the supplier docs; Country of
Origin from the AI compare step; Notify Party from the customer list;
Delivery and Payment Term static. The item block (material, batches,
manufacturer, weights) comes from the COA/docs; company registration/bank
details stay static; totals (in words + number for invoice, weights for
packing) are computed sums across items. The stamp is static.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional

import docx
from docx.oxml.ns import qn
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from forms_docx_utils import (
    clone_row_for_items,
    replace_cell_text,
    replace_regex_span,
    set_paragraph_full_text,
    sum_weight_texts,
    _unique_row_cells,
)
from forms_vigorous import _customer_lines, _fmt_money, _item_total
import num2words_en

FORM_DIR = Path(__file__).resolve().parent.parent / "forms" / "sedanah_jordan"


def _textbox_paragraph_groups(doc) -> List[List[Paragraph]]:
    """Both Sedanah templates put the 'INV No./Dated' and 'To:' fields inside
    floating text boxes (anchored drawings), not the regular body paragraphs
    or table cells -- python-docx's normal doc.paragraphs/table.cells never
    see them. Word/LibreOffice duplicate each floating box into two parallel
    XML representations (DrawingML + a VML fallback) with identical text, so
    every group below appears twice; both copies must be updated to keep
    them in sync regardless of which one a given viewer renders."""
    body = doc.element.body
    groups = []
    for tb in body.findall(".//" + qn("w:txbxContent")):
        paras = [Paragraph(p, None) for p in tb.findall(qn("w:p"))]
        groups.append(paras)
    return groups


def _fill_invoice_no_textboxes(doc, invoice_no: str, invoice_date: str) -> None:
    for group in _textbox_paragraph_groups(doc):
        if group and group[0].text.strip().startswith("INV No"):
            replace_regex_span(group[0], r"(INV No\.:\s*)(.*)$", invoice_no or "N/A", group=2)
            if len(group) > 1:
                replace_regex_span(group[1], r"(Dated:\s*)(.*)$", invoice_date or "N/A", group=2)


def _fill_to_textboxes(doc, customer: dict) -> None:
    lines = _customer_lines(customer)
    for group in _textbox_paragraph_groups(doc):
        if group and group[0].text.strip() == "To:":
            addr_paras = group[1:]
            for i, p in enumerate(addr_paras):
                if i < len(lines):
                    text = lines[i] if i < len(lines) - 1 or len(lines) <= len(addr_paras) else ", ".join(lines[i:])
                else:
                    text = ""
                set_paragraph_full_text(p, text)


def _find_cell(table: Table, label: str) -> Optional[_Cell]:
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            if label in cell.text:
                return cell
    return None


def _replace_second_line(cell: Optional[_Cell], new_value: str) -> None:
    """For label+value cells that come in either of two template styles --
    'Port of Discharge\\nALEX' (newline-separated) or 'Country of Origin :
    INDIA' (colon-separated on one line) -- keep the label, replace the
    value."""
    if cell is None or not cell.paragraphs:
        return
    text = cell.text
    if "\n" in text:
        label = text.split("\n", 1)[0]
        replace_cell_text(cell, f"{label}\n{new_value}")
    elif ":" in text:
        label = text.rsplit(":", 1)[0]
        replace_cell_text(cell, f"{label}:  {new_value}")
    else:
        replace_cell_text(cell, new_value)


def _find_value_cell_below_label(table: Table, exact_label: str) -> Optional[_Cell]:
    """For header cells where the label ('Carriage by', 'Port of Loading') sits
    alone in one row and the value sits in the SAME column of the next row
    (two separate, unmerged cells) -- find and return that value cell."""
    rows = list(table.rows)
    for idx, row in enumerate(rows):
        cells = _unique_row_cells(row)
        for col, cell in enumerate(cells):
            if cell.text.strip() == exact_label and idx + 1 < len(rows):
                below = _unique_row_cells(rows[idx + 1])
                if col < len(below):
                    return below[col]
    return None


def _fill_header(table: Table, customer: dict, shipment) -> None:
    notify = _find_cell(table, "Notify Party")
    if notify is not None:
        replace_cell_text(notify, "Notify Party:\n" + "\n".join(_customer_lines(customer)))

    _replace_second_line(_find_cell(table, "Port of Discharge"), shipment.port_of_discharge or "N/A")
    _replace_second_line(_find_cell(table, "Port of Final Destination"), shipment.port_of_final_destination or "N/A")
    _replace_second_line(_find_cell(table, "Country of Origin"), (shipment.origin or "N/A").upper())
    _replace_second_line(_find_cell(table, "Country of Final Destination"), shipment.country_of_final_destination or "N/A")

    loading_cell = _find_value_cell_below_label(table, "Port of Loading")
    if loading_cell is not None:
        replace_cell_text(loading_cell, shipment.port_of_loading or "N/A")

    carriage_cell = _find_value_cell_below_label(table, "Carriage by")
    if carriage_cell is not None:
        replace_cell_text(carriage_cell, shipment.carriage_by or "N/A")


def _item_block_lines(item: dict) -> List[str]:
    lines = [item.get("material_name", "")]
    for b in item.get("batches") or [{"batch_no": "", "manufacturing_date": "", "expiry_date": ""}]:
        lines.append(f"BATCH No: {b.get('batch_no', '')} ")
        lines.append(f"MFG Date: {b.get('manufacturing_date', '')}         EXP Date: {b.get('expiry_date', '')}")
    lines.append(f"MANUFACTURED BY: {item.get('manufacturer', '') or 'N/A'}")
    lines.append(f"GR.WT. {item.get('gross_weight_text', '')}           NT.WT. {item.get('net_weight_text', '')}")
    return lines


def _fill_invoice(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    items = shipment.items or []
    table0 = doc.tables[0]  # "From:" -- static, untouched
    table1 = doc.tables[1]

    _fill_invoice_no_textboxes(doc, shipment.invoice_no, shipment.invoice_date)
    _fill_to_textboxes(doc, customer)
    _fill_header(table1, customer, shipment)

    # Static company/bank boilerplate tail from the template's item cell,
    # kept exactly as-is and appended once (after the first item's block).
    item_row = None
    for row in table1.rows:
        cells = _unique_row_cells(row)
        if len(cells) == 5 and "BATCH No" in cells[1].text:
            item_row = row
            break

    static_tail = ""
    if item_row is not None:
        text = _unique_row_cells(item_row)[1].text
        if "CARGOX ID" in text:
            static_tail = "\n\n" + text[text.index("CARGOX ID"):]

    def render_item_row(idx_item):
        idx, item = idx_item
        block = "\n".join(_item_block_lines(item))
        if idx == 0:
            block += static_tail
        currency = item.get("currency") or "USD"
        return [
            f"{idx + 1}.",
            block,
            item.get("quantity_text", ""),
            f"{currency} {item.get('unit_price', 0):g}/{item.get('unit', '')}",
            f"{currency} {_fmt_money(_item_total(item))}",
        ]

    if item_row is not None:
        indexed = list(enumerate(items)) or [(0, {})]
        clone_row_for_items(item_row, indexed, render_item_row)

    grand_total = sum(_item_total(i) for i in items)
    currency = (items[0].get("currency") if items else "") or "USD"
    for row in table1.rows:
        cells = _unique_row_cells(row)
        if len(cells) == 3 and cells[1].text.strip().upper() == "TOTAL:":
            replace_cell_text(cells[0], f"SAY: {num2words_en.amount_to_words(grand_total, currency)}")
            replace_cell_text(cells[2], f"{currency} {_fmt_money(grand_total)}")
            break

    origin = (shipment.origin or "").upper()
    for p in doc.paragraphs:
        if "exclusively" in p.text:
            replace_regex_span(p, r"\(([A-Za-z ]+)\)\s*$", origin, group=1)
            break

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_packing(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    items = shipment.items or []
    table1 = doc.tables[1]

    _fill_invoice_no_textboxes(doc, shipment.invoice_no, shipment.invoice_date)
    _fill_to_textboxes(doc, customer)
    _fill_header(table1, customer, shipment)

    item_row = None
    for row in table1.rows:
        cells = _unique_row_cells(row)
        if len(cells) == 2 and "BATCH No" in cells[1].text:
            item_row = row
            break

    static_tail = ""
    if item_row is not None:
        text = _unique_row_cells(item_row)[1].text
        if "CARGOX ID" in text:
            tail_start = text.index("CARGOX ID")
            tail_end = text.index("QUANTITY:") if "QUANTITY:" in text else len(text)
            static_tail = "\n\n" + text[tail_start:tail_end].rstrip("\n")

    def render_item_row(idx_item):
        idx, item = idx_item
        lines = _item_block_lines(item)[:-1]  # drop the GR/NT weight line -- rebuilt below
        block = "\n".join(lines)
        if idx == 0:
            block += static_tail
        block += (
            f"\n\nQUANTITY: {item.get('quantity_text', '')}"
            f"\n{item.get('packaging_description', '')}"
            f"\nGR.WT. {item.get('gross_weight_text', '')}           NT.WT. {item.get('net_weight_text', '')}"
        )
        return [f"\n{idx + 1}.", block]

    if item_row is not None:
        indexed = list(enumerate(items)) or [(0, {})]
        clone_row_for_items(item_row, indexed, render_item_row)

    gross_total = sum_weight_texts([i.get("gross_weight_text", "") for i in items])
    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )
    for row in table1.rows:
        cells = _unique_row_cells(row)
        if len(cells) == 2 and cells[1].text.strip().upper().startswith("TOTAL GR"):
            replace_cell_text(cells[1], f"TOTAL GR.WT. {gross_total}-   TOTAL NT.WT. {net_total} - TOTAL {package_desc}")
            break

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def generate(customer: dict, shipment, output_dir: Path) -> Dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    results = {}
    results["invoice"] = output_dir / "invoice.docx"
    _fill_invoice(FORM_DIR / "invoice.docx", results["invoice"], customer, shipment)

    results["packing"] = output_dir / "packing_list.docx"
    _fill_packing(FORM_DIR / "packing.docx", results["packing"], customer, shipment)

    return results
