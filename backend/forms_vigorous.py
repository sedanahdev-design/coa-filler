"""Fill logic for the 'Vigorous Form' (invoice, packing list, certificate of
origin) -- the first of the 5 customer document-forms.

Field-mapping rules (as specified by the user, form-by-form):

Invoice: No/Date from docs; customer block auto-filled from the chosen
customer; item lines (one per shipment item) with Total = qty * unit price,
computed automatically; batch/mfg/exp lines per item from docs; VALUE TOTAL
= sum of item totals; SAY line = that total in words; Manufacturer from COA;
Supplier/Declarations/Stamp untouched except the Origin value (and the
Origin mentioned inside the declaration parentheses), both from the AI
origin-comparison step.

Packing list: customer block from customer list; invoice No/Date from docs;
one packaging row per item (material/packaging/weights from docs); footer
totals = sum across items; stamp untouched.

COO: consignee block from customer list; package count (box 6) and material
name (box 7, one line only) and origin (box 8) dynamic; invoice no/date
(box 10) from docs; net weight (box 9) from docs; everything else static.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Dict, List

import docx

from forms_docx_utils import (
    clone_item_with_sublines_for_items,
    clone_paragraph_for_items,
    clone_row_for_items,
    fill_multiline_block,
    find_paragraph,
    replace_cell_text,
    replace_paragraph_span,
    replace_regex_span,
    set_paragraph_full_text,
    sum_weight_texts,
)
import num2words_en

FORM_DIR = Path(__file__).resolve().parent.parent / "forms" / "vigorous"


def _item_total(item: dict) -> float:
    return float(item.get("quantity_value") or 0) * float(item.get("unit_price") or 0)


def _flatten_batches(items: List[dict]) -> List[dict]:
    """Flatten each item's `batches` array into a single list of batch dicts
    (material_name carried along for context), preserving item order. Items
    with no batches listed fall back to one blank placeholder batch so a line
    still gets rendered."""
    flat = []
    for item in items:
        batches = item.get("batches") or [
            {"batch_no": "", "manufacturing_date": "", "expiry_date": "", "retest_date": ""}
        ]
        for b in batches:
            flat.append({**b, "material_name": item.get("material_name", "")})
    return flat


def _customer_lines(customer: dict) -> List[str]:
    text = customer.get("full_text") or customer.get("name") or ""
    return [l for l in text.splitlines() if l.strip()]


def _fmt_money(value: float) -> str:
    return f"{value:,.2f}".rstrip("0").rstrip(".") if value == int(value) else f"{value:,.2f}"


def _fill_invoice(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs

    # No / Dated -- leave truly blank (not "N/A") when the source docs didn't
    # have a clearly-labeled invoice number/date, rather than filling in a
    # placeholder that reads as if it were a real, checked value.
    replace_regex_span(
        paras[2],
        r"No:\s*XX\s*Dated:\s*XX/XX/2026",
        f"No: {shipment.invoice_no} Dated: {shipment.invoice_date}",
    )

    # Customer block (paragraphs 4-6: name / address line / phone line)
    fill_multiline_block(paras[4:7], _customer_lines(customer))

    items = shipment.items or []

    # Item line (paragraph 9) cloned once per item
    def render_item_line(item: dict) -> str:
        total = _item_total(item)
        name = item.get("material_name", "")
        qty = item.get("quantity_text", "")
        price = item.get("unit_price", 0)
        currency = item.get("currency") or "USD"
        unit = item.get("unit", "")
        return (
            f"{name:<36}{qty:<25}{currency} {price:g}/{unit}"
            f"{'':<28}{currency} {_fmt_money(total)}"
        )

    # Batch/Mfg/Exp line (paragraph 11): one per batch of THAT item (a
    # material is often split across several batch numbers). Each item's
    # batch line(s) are rendered directly under that item's own line, the
    # same way the supplier invoice lays them out -- previously all item
    # lines were cloned first and every batch was flattened together at the
    # end, so batches no longer sat under the item they belong to.
    def render_batch_line(batch: dict, item: dict = None) -> str:
        # Prefer a genuine expiry date; a lot of real CoAs (reference
        # standards, some bulk APIs) state only a Retest/Re-qualification
        # date instead of a hard expiry -- confirmed real gap: with no
        # fallback here, this line's date came out blank for every one of
        # those, which looked like the expiry date was "never" being
        # picked up at all. Falling back to the retest date (relabeled)
        # means the line always shows *some* usable date when the source
        # states one, rather than silently leaving it empty.
        expiry = batch.get("expiry_date", "")
        if expiry:
            date_label, date_value = "Exp. Date:", expiry
        else:
            retest = batch.get("retest_date", "")
            date_label, date_value = ("Retest Date:", retest) if retest else ("Exp. Date:", "")
        return (
            f"                 Batch NO: {batch.get('batch_no', '')}"
            f"             Mfg. Date: {batch.get('manufacturing_date', '')}"
            f"                 {date_label} {date_value}"
        )

    clone_item_with_sublines_for_items(
        paras[9],
        paras[11],
        items,
        render_item_line,
        lambda item: item.get("batches") or [],
        render_batch_line,
    )

    grand_total = sum(_item_total(i) for i in items)

    # VALUE TOTAL line -- only trailing "USD n" number changes
    p14 = paras[14]
    replace_regex_span(p14, r"(USD|[A-Z]{3})\s*[\d,]+(?:\.\d+)?\s*$", f"USD {_fmt_money(grand_total)}")

    # SAY line -- amount in words
    currency = (items[0].get("currency") if items else "") or "USD"
    set_paragraph_full_text(paras[15], "SAY " + num2words_en.amount_to_words(grand_total, currency))

    # Manufacturer
    manufacturers = sorted({i.get("manufacturer", "") for i in items if i.get("manufacturer")})
    replace_regex_span(paras[16], r"(Manufacturer\s+)(.*)$", ", ".join(manufacturers) or "N/A", group=2)

    # Origin: line
    origin = (shipment.origin or "").upper()
    replace_regex_span(paras[18], r"(Origin:\s*)(\S.*?)\s*$", origin, group=2)

    # Origin inside the declaration parentheses
    replace_regex_span(paras[21], r"\(([A-Za-z ]+)\)\s*$", origin, group=1)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_packing(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs

    fill_multiline_block(paras[5:8], _customer_lines(customer))

    replace_regex_span(
        paras[8],
        r"No:\s*XX\s*Dated:\s*XX/XX/2026",
        f"No: {shipment.invoice_no} Dated: {shipment.invoice_date}",
    )

    table = doc.tables[0]
    items = shipment.items or []

    def render_item_row(item: dict) -> List[str]:
        marks = "N/M"
        desc = (
            f"{item.get('material_name', '')}\n\n"
            f"{item.get('packaging_description', '')}\n"
            f"GR.WT. {item.get('gross_weight_text', '')}.\n"
            f"NT.WT. {item.get('net_weight_text', '')}.\n"
        )
        return [marks, desc]

    if items:
        clone_row_for_items(table.rows[1], items, render_item_row)
    else:
        for cell in table.rows[1].cells:
            replace_cell_text(cell, "")

    # Footer totals row is now the last row (after any cloned item rows)
    totals_row = table.rows[len(table.rows) - 1]
    gross_total = sum_weight_texts([i.get("gross_weight_text", "") for i in items])
    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )
    totals_cells = [c for c in totals_row.cells]
    seen = set()
    unique_cells = []
    for c in totals_cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        unique_cells.append(c)
    if len(unique_cells) >= 2:
        replace_cell_text(unique_cells[1], f"TOTAL GR. WT {gross_total}   -   TOTAL NT.WT. {net_total}\nTOTAL {package_desc}.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _cell_paragraphs(table, row_idx: int, col_idx: int):
    from forms_docx_utils import _unique_row_cells

    cells = _unique_row_cells(table.rows[row_idx])
    return cells[col_idx].paragraphs, cells[col_idx]


def _fill_coo(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    table = doc.tables[0]
    items = shipment.items or []

    # Box 2: consignee block
    consignee_paras, _ = _cell_paragraphs(table, 1, 0)
    fill_multiline_block(consignee_paras[4:8], _customer_lines(customer))

    # Box 6: marks & no. of packages
    marks_paras, _ = _cell_paragraphs(table, 4, 1)
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )
    if len(marks_paras) > 3:
        set_paragraph_full_text(marks_paras[3], package_desc or "N/A")

    # Box 7: material name (line 3 of the description only)
    goods_paras, _ = _cell_paragraphs(table, 4, 2)
    material_names = ", ".join(sorted({i.get("material_name", "") for i in items if i.get("material_name")}))
    if len(goods_paras) > 7:
        set_paragraph_full_text(goods_paras[7], material_names or "N/A")

    # Box 8: origin criteria
    origin_paras, _ = _cell_paragraphs(table, 4, 3)
    origin = (shipment.origin or "").upper()
    if len(origin_paras) > 3:
        set_paragraph_full_text(origin_paras[3], f"    {origin}")

    # Box 9: net weight
    weight_paras, _ = _cell_paragraphs(table, 4, 4)
    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    if len(weight_paras) > 4:
        set_paragraph_full_text(weight_paras[4], net_total or "N/A")

    # Box 10: invoice no / date -- leave blank rather than "N/A" if not found
    inv_paras, _ = _cell_paragraphs(table, 4, 5)
    if len(inv_paras) > 4:
        set_paragraph_full_text(inv_paras[2], f" {shipment.invoice_no} " if shipment.invoice_no else "")
        set_paragraph_full_text(inv_paras[4], shipment.invoice_date)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def generate(customer: dict, shipment, output_dir: Path) -> Dict[str, Path]:
    """Generate the 3 Vigorous-form documents. `shipment` is a
    forms_extract.ShipmentData (or any object/duck-type with the same
    attributes: invoice_no, invoice_date, origin, total_package_description,
    items). Returns {"invoice": path, "packing": path, "coo": path}."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    results = {}
    results["invoice"] = output_dir / "invoice.docx"
    _fill_invoice(FORM_DIR / "invoice.docx", results["invoice"], customer, shipment)

    results["packing"] = output_dir / "packing_list.docx"
    _fill_packing(FORM_DIR / "packing.docx", results["packing"], customer, shipment)

    results["coo"] = output_dir / "certificate_of_origin.docx"
    _fill_coo(FORM_DIR / "coo.docx", results["coo"], customer, shipment)

    return results
