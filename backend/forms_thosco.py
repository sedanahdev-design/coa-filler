"""Fill logic for the 'Thosco' form (invoice, packing/weight list, COO).

Field mapping (per user spec): Buyer & Consignee from the customer list;
invoice No./Date from docs; MANUFACTURER from COA; ORIGIN from the AI
compare step; DELIVERY TERMS / PORT / TERMS OF PAYMENT static; table/total
follow the same computed pattern as the other forms (quantity*price, summed
for the total, in words). A single product commonly ships across several
batch numbers, so the batch lines are cloned once per batch (flattened
across items), while the product/quantity line is cloned once per item.

Packing: product/packed-in/weights per item from COA/docs; TOTAL PACKAGES =
sum across all items; logo/company static.

COO: exporter block static; consignee from customer list; origin from AI;
delivery terms static; packaging description + invoice no from docs;
manufacturer name from COA.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List

import docx

from forms_docx_utils import (
    clone_paragraph_for_items,
    clone_paragraph_group_for_items,
    fill_multiline_block,
    replace_cell_text,
    replace_paragraph_span,
    replace_regex_span,
    set_paragraph_full_text,
    sum_weight_texts,
    _unique_row_cells,
)
from forms_vigorous import _customer_lines, _fmt_money, _flatten_batches, _item_total
import num2words_en

FORM_DIR = Path(__file__).resolve().parent.parent / "forms" / "thosco"

_INDENT = "                                                 "


def _fill_no_dated(paragraph, invoice_no: str, invoice_date: str) -> None:
    replace_regex_span(
        paragraph,
        r"NO\.\s*\S+\s*DATE:\s*.+$",
        f"NO. {invoice_no or 'N/A'} DATE: {invoice_date or 'N/A'}",
    )


def _fill_customer_block(para_label, para2, para3, customer: dict) -> None:
    lines = _customer_lines(customer)
    replace_regex_span(para_label, r"(Buyer & Consignee Party:\s*)(.*)$", lines[0] if lines else "N/A", group=2)
    set_paragraph_full_text(para2, _INDENT + lines[1] if len(lines) > 1 else "")
    set_paragraph_full_text(para3, _INDENT + ", ".join(lines[2:]) if len(lines) > 2 else "")


def _fill_invoice(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs
    items = shipment.items or []

    _fill_no_dated(paras[3], shipment.invoice_no, shipment.invoice_date)
    _fill_customer_block(paras[4], paras[5], paras[6], customer)

    def render_item_line(item: dict) -> str:
        total = _item_total(item)
        name = item.get("material_name", "")
        qty = item.get("quantity_text", "")
        price = item.get("unit_price", 0)
        currency = item.get("currency") or "USD"
        unit = item.get("unit", "")
        return (
            f"{name:<58}{qty:<32}{currency} {price:g}/{unit}"
            f"{'':<24}{currency} {_fmt_money(total)}"
        )

    clone_paragraph_for_items(paras[12], items, render_item_line)

    flat_batches = _flatten_batches(items)

    def render_batch_line(batch: dict) -> str:
        return (
            f"                   BATCH NO.:   {batch.get('batch_no', '')}"
            f"              MFG. DATE: {batch.get('manufacturing_date', '')}"
            f"          /           EXP. DATE: {batch.get('expiry_date', '')}    "
        )

    clone_paragraph_for_items(paras[14], flat_batches, render_batch_line)

    # The template ships with 5 extra sample batch lines (paragraphs 15-19)
    # beyond the first (paragraph 14) -- remove them now that our own batch
    # lines have been cloned in after paragraph 14, otherwise the sample
    # data leaks through as stray extra lines.
    for p in paras[15:20]:
        el = p._p
        if el.getparent() is not None:
            el.getparent().remove(el)

    grand_total = sum(_item_total(i) for i in items)
    currency = (items[0].get("currency") if items else "") or "USD"
    replace_regex_span(
        paras[22],
        r"(TOTAL AMOUNT:\s*)(.*)$",
        f"({currency} {_fmt_money(grand_total)})  SAY {num2words_en.amount_to_words(grand_total, currency)}",
        group=2,
    )

    manufacturers = sorted({i.get("manufacturer", "") for i in items if i.get("manufacturer")})
    replace_regex_span(paras[23], r"(MANUFACTURER\s*:\s*)(.*)$", ", ".join(manufacturers) or "N/A", group=2)

    origin = (shipment.origin or "").upper()
    replace_regex_span(paras[24], r"(ORIGIN\s*:\s*)(\S.*?)\s*$", origin, group=2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_packing(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs
    items = shipment.items or []

    _fill_no_dated(paras[3], shipment.invoice_no, shipment.invoice_date)
    _fill_customer_block(paras[5], paras[6], paras[7], customer)

    def render_product(item: dict) -> str:
        return f"PRODUCT:                               {item.get('material_name', '')}"

    def render_packed_in(item: dict) -> str:
        return f"PACKED IN:                             {item.get('packaging_description', '')}"

    clone_paragraph_group_for_items(
        [paras[11], paras[12]], items, [render_product, render_packed_in]
    )

    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    gross_total = sum_weight_texts([i.get("gross_weight_text", "") for i in items])
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )

    set_paragraph_full_text(paras[14], f"TOTAL NET WIGHT:               {net_total}")
    set_paragraph_full_text(paras[15], f"TOTAL GROSS WEIGHT:        {gross_total}")
    set_paragraph_full_text(paras[16], f"TOTAL PACKAGES:                 {package_desc}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_coo(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    table = doc.tables[0]
    items = shipment.items or []

    # Consignee (row 4/5 share the same vertically-merged cell)
    consignee_cell = _unique_row_cells(table.rows[4])[0]
    replace_cell_text(consignee_cell, "\n\n" + "\n".join(_customer_lines(customer)) + "\n")

    # Origin (row 5, col 1, second paragraph)
    origin_cell = _unique_row_cells(table.rows[5])[1]
    origin = (shipment.origin or "").upper()
    if len(origin_cell.paragraphs) > 1:
        set_paragraph_full_text(origin_cell.paragraphs[1], f" \n {origin}   ")

    # Goods description + invoice + manufacturer (row 8, col 0)
    goods_cell = _unique_row_cells(table.rows[8])[0]
    gp = goods_cell.paragraphs
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )
    material_names = ", ".join(sorted({i.get("material_name", "") for i in items if i.get("material_name")}))
    manufacturers = sorted({i.get("manufacturer", "") for i in items if i.get("manufacturer")})
    if len(gp) > 8:
        set_paragraph_full_text(gp[2], f"{package_desc} OF {material_names}".ljust(90))
        set_paragraph_full_text(gp[5], f"INVOICE NO: {shipment.invoice_no or 'N/A'} DATED :{shipment.invoice_date or 'N/A'}")
        set_paragraph_full_text(gp[6], f"ORIGIN OF THE GOODS IS {origin}")
        set_paragraph_full_text(gp[7], f"MANUFACTURER NAME: {', '.join(manufacturers) or 'N/A'}")
        set_paragraph_full_text(gp[8], f" {origin}")

    # Weight (row 8, col 1)
    weight_cell = _unique_row_cells(table.rows[8])[1]
    wp = weight_cell.paragraphs
    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    if len(wp) > 3:
        set_paragraph_full_text(wp[3], f"       NET: {net_total}  ")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def generate(customer: dict, shipment, output_dir: Path) -> Dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    results = {}
    results["invoice"] = output_dir / "invoice.docx"
    _fill_invoice(FORM_DIR / "invoice.docx", results["invoice"], customer, shipment)

    results["packing"] = output_dir / "packing_weight_list.docx"
    _fill_packing(FORM_DIR / "packing.docx", results["packing"], customer, shipment)

    results["coo"] = output_dir / "certificate_of_origin.docx"
    _fill_coo(FORM_DIR / "coo.docx", results["coo"], customer, shipment)

    return results
