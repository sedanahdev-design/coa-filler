"""Fill logic for the 'Biesterfeld Germany' form (invoice, packing, COO).

Invoice: manufacturer/buyer address block from the customer list; Date and
Invoice No. from docs; MANUFACTURER from COA; SUPPLIER static; ORIGIN from
the AI compare step (appears twice: the ORIGIN: field and inside the
declaration text); delivery/port/payment terms static; TOTAL NET WIGHT
summed across all materials from docs.

Packing and COO reuse the same field-mapping pattern as Thosco (per the
user: "the same of the packing/COO in the previous form the same way").
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import docx

from forms_docx_utils import (
    clone_paragraph_for_items,
    clone_paragraph_group_for_items,
    replace_cell_text,
    replace_regex_span,
    set_paragraph_full_text,
    sum_weight_texts,
    _unique_row_cells,
)
from forms_vigorous import _customer_lines, _fmt_money, _flatten_batches, _item_total
import num2words_en

FORM_DIR = Path(__file__).resolve().parent.parent / "forms" / "biesterfeld_germany"


def _fill_customer_block(paras, start_idx: int, customer: dict, indent: str = "      ") -> None:
    lines = _customer_lines(customer)
    for i in range(4):
        text = indent + lines[i] if i < len(lines) else ""
        set_paragraph_full_text(paras[start_idx + i], text)


def _fill_invoice(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs
    items = shipment.items or []

    _fill_customer_block(paras, 7, customer, indent="")

    replace_regex_span(paras[11], r"Date:\s*.+$", f"Date: {shipment.invoice_date or 'N/A'}")
    replace_regex_span(paras[13], r"(INVOICE NO\.:\s*)(.*)$", shipment.invoice_no or "N/A", group=2)

    def render_item_line(item: dict) -> str:
        total = _item_total(item)
        name = item.get("material_name", "")
        qty = item.get("quantity_text", "")
        price = item.get("unit_price", 0)
        currency = item.get("currency") or "USD"
        return f"{name:<36}{qty:<32}{currency} {price:g}{'':<38}{currency} {_fmt_money(total)}"

    clone_paragraph_for_items(paras[18], items, render_item_line)

    flat_batches = _flatten_batches(items)

    def render_batch_line(batch: dict) -> str:
        return (
            f"           Batch no.:  {batch.get('batch_no', '')}"
            f"               Mfg. Date: {batch.get('manufacturing_date', '')}"
            f"   /           Ret. Date: {batch.get('expiry_date', '')}"
        )

    clone_paragraph_for_items(paras[20], flat_batches, render_batch_line)

    grand_total = sum(_item_total(i) for i in items)
    currency = (items[0].get("currency") if items else "") or "USD"
    currency_symbol = {"USD": "$", "EUR": "€", "GBP": "£"}.get(currency, currency)
    replace_regex_span(
        paras[24],
        r"(TOTAL AMOUNT:\s*)(.*)$",
        f"({_fmt_money(grand_total)}{currency_symbol})   SAY {num2words_en.amount_to_words(grand_total, currency)}",
        group=2,
    )

    manufacturers = sorted({i.get("manufacturer", "") for i in items if i.get("manufacturer")})
    replace_regex_span(paras[25], r"(MANUFACTURER:\s*)(.*)$", ", ".join(manufacturers) or "N/A", group=2)

    origin = (shipment.origin or "").upper()
    replace_regex_span(paras[27], r"(ORIGIN:\s*)(\S.*?)\s*$", origin, group=2)

    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    replace_regex_span(paras[31], r"(TOTAL NET WIGHT:\s*)(.*)$", net_total or "N/A", group=2)

    replace_regex_span(paras[34], r"(exclusively from\s*)(.*)$", origin, group=2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_packing(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    paras = doc.paragraphs
    items = shipment.items or []

    _fill_customer_block(paras, 8, customer, indent="      ")

    replace_regex_span(
        paras[18],
        r"Invoice Number:\s*\S+\s*Date:\s*.+$",
        f"Invoice Number:  {shipment.invoice_no or 'N/A'}                                               Date: {shipment.invoice_date or 'N/A'}",
    )

    def render_product(item: dict) -> str:
        return f"PRODUCT:                            {item.get('material_name', '')}"

    def render_packed_in(item: dict) -> str:
        return f"PACKED IN:                           {item.get('packaging_description', '')}"

    clone_paragraph_group_for_items([paras[22], paras[23]], items, [render_product, render_packed_in])

    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    gross_total = sum_weight_texts([i.get("gross_weight_text", "") for i in items])
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )

    set_paragraph_full_text(paras[24], f"TOTAL NET WIGHT:               {net_total}")
    set_paragraph_full_text(paras[25], f"TOTAL GROSS WEIGHT:       {gross_total}")
    set_paragraph_full_text(paras[26], f"TOTAL PACKAGES:               {package_desc}")

    origin = (shipment.origin or "").upper()
    replace_regex_span(paras[27], r"(ORIGIN:\s*)(\S.*?)\s*$", origin, group=2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _fill_coo(path: Path, out_path: Path, customer: dict, shipment) -> None:
    doc = docx.Document(str(path))
    table = doc.tables[0]
    items = shipment.items or []

    consignee_cell = _unique_row_cells(table.rows[4])[0]
    lines = _customer_lines(customer)
    replace_cell_text(consignee_cell, "\n\n" + "\n".join(lines) + "  ")

    origin_cell = _unique_row_cells(table.rows[5])[1]
    origin = (shipment.origin or "").upper()
    if len(origin_cell.paragraphs) > 1:
        set_paragraph_full_text(origin_cell.paragraphs[1], f" \n {origin}   ")

    goods_cell = _unique_row_cells(table.rows[8])[0]
    gp = goods_cell.paragraphs
    package_desc = shipment.total_package_description or sum_weight_texts(
        [i.get("package_count_text", "") for i in items]
    )
    material_names = ", ".join(sorted({i.get("material_name", "") for i in items if i.get("material_name")}))
    manufacturers = sorted({i.get("manufacturer", "") for i in items if i.get("manufacturer")})
    if len(gp) > 9:
        set_paragraph_full_text(gp[2], f"{package_desc} OF {material_names}")
        set_paragraph_full_text(gp[5], f"INVOICE NO: {shipment.invoice_no or 'N/A'} DATED:  {shipment.invoice_date or 'N/A'}")
        set_paragraph_full_text(gp[6], f"ORIGIN OF THE GOODS IS {origin}")
        set_paragraph_full_text(gp[8], ", ".join(manufacturers) or "N/A")
        set_paragraph_full_text(gp[9], origin)

    weight_cell = _unique_row_cells(table.rows[8])[1]
    wp = weight_cell.paragraphs
    net_total = sum_weight_texts([i.get("net_weight_text", "") for i in items])
    if len(wp) > 3:
        set_paragraph_full_text(wp[3], f"       NET: {net_total} ")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def generate(customer: dict, shipment, output_dir: Path) -> Dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    results = {}
    results["invoice"] = output_dir / "invoice.docx"
    _fill_invoice(FORM_DIR / "invoice.docx", results["invoice"], customer, shipment)

    results["packing"] = output_dir / "packing_list.docx"
    _fill_packing(FORM_DIR / "packing.docx", results["packing"], customer, shipment)

    results["coo"] = output_dir / "certificate_of_origin.docx"
    _fill_coo(FORM_DIR / "coo.docx", results["coo"], customer, shipment)

    return results
