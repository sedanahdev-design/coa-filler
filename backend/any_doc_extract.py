"""
Generic "read any document" extraction, used by the comparison tool and the
forms pipeline. Supports PDF, common image formats, and .docx -- returns a
uniform structure (raw text + zero or more page/whole-document images as
base64 data URLs) that ai_extract-style OpenAI calls can consume directly.

No network access.
"""
from __future__ import annotations

import base64
import dataclasses
import mimetypes
from pathlib import Path
from typing import List

import docx as docx_lib

import pdf_extract

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}
DOCX_EXTS = {".docx"}
PDF_EXTS = {".pdf"}


@dataclasses.dataclass
class DocContent:
    filename: str
    kind: str  # "pdf" | "image" | "docx"
    text: str
    image_data_urls: List[str]
    text_reliable: bool = True


def _image_to_data_url(path: str) -> str:
    mime, _ = mimetypes.guess_type(path)
    mime = mime or "image/png"
    with open(path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("ascii")
    return f"data:{mime};base64,{b64}"


def _extract_docx_text(path: str) -> str:
    doc = docx_lib.Document(path)
    chunks = []
    for p in doc.paragraphs:
        if p.text.strip():
            chunks.append(p.text)

    def walk_table(table):
        for row in table.rows:
            seen = set()
            cells_text = []
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                cells_text.append(cell.text.strip())
                for nested in cell.tables:
                    walk_table(nested)
            if any(cells_text):
                chunks.append(" | ".join(cells_text))

    for table in doc.tables:
        walk_table(table)

    return "\n".join(chunks)


def read_document(path: str, filename: str = None, pdf_dpi: int = 200) -> DocContent:
    filename = filename or Path(path).name
    ext = Path(filename).suffix.lower()

    if ext in PDF_EXTS:
        extraction = pdf_extract.extract_all(path, dpi=pdf_dpi)
        return DocContent(
            filename=filename,
            kind="pdf",
            text=extraction.text,
            image_data_urls=[p.as_data_url() for p in extraction.pages],
            text_reliable=extraction.text_reliable,
        )

    if ext in IMAGE_EXTS:
        return DocContent(
            filename=filename,
            kind="image",
            text="",
            image_data_urls=[_image_to_data_url(path)],
        )

    if ext in DOCX_EXTS:
        return DocContent(
            filename=filename,
            kind="docx",
            text=_extract_docx_text(path),
            image_data_urls=[],
        )

    raise ValueError(f"Unsupported file type: {ext or filename}")
